import io
import os
from datetime import date, datetime
import requests
import openpyxl

TOKEN = os.environ["BOT_TOKEN"]
CHAT_ID = os.environ["AUTO_CHAT_ID"]
SPREADSHEET_ID = "1UmJKEHIy7vYc3matpjZ9ZC7NqiVY3QHETa4f4DotAlY"
SUBJECT_COL = 11
ROOM_COL = 12
DATE_COL = 1
PAIR_COL = 3
TIME_COL = 4


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def day_name(d):
    return ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"][d.weekday()]


def get_schedule(target_date):
    url = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=xlsx"
    response = requests.get(url, timeout=30)
    response.raise_for_status()
    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    result = []

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        current_date = None
        for row in range(1, ws.max_row + 1):
            parsed = normalize_date(ws.cell(row=row, column=DATE_COL).value)
            if parsed:
                current_date = parsed
            if current_date != target_date:
                continue

            pair_value = ws.cell(row=row, column=PAIR_COL).value
            lesson_time = clean(ws.cell(row=row, column=TIME_COL).value)
            subject = clean(ws.cell(row=row, column=SUBJECT_COL).value)
            room = clean(ws.cell(row=row, column=ROOM_COL).value)

            if pair_value is not None and subject:
                try:
                    pair = int(float(pair_value))
                except (ValueError, TypeError):
                    pair = None
                teacher = None
                if row + 1 <= ws.max_row:
                    possible_teacher = clean(ws.cell(row=row + 1, column=SUBJECT_COL).value)
                    if possible_teacher and possible_teacher != subject:
                        teacher = possible_teacher
                result.append({"pair": pair, "time": lesson_time, "subject": subject, "teacher": teacher, "room": room})
            elif lesson_time and subject:
                if "преподаватель:" not in subject.lower():
                    result.append({"pair": None, "time": lesson_time, "subject": subject, "teacher": None, "room": room})

    unique, seen = [], set()
    for item in result:
        key = (item["pair"], item["time"], item["subject"], item["room"])
        if key not in seen:
            seen.add(key)
            unique.append(item)
    unique.sort(key=lambda x: (x["pair"] is None, x["pair"] if x["pair"] is not None else 999))
    return unique


def format_day(target_date, lessons):
    text = f"📅 <b>{day_name(target_date)}</b>\n📆 {target_date:%d.%m.%Y}\n🎓 <b>Группа 451</b>\n\n"
    if not lessons:
        return text + "😎 <b>Пар нет.</b> Можно отдыхать!"
    for lesson in lessons:
        text += f"🔔 <b>{lesson['pair']} пара</b>" if lesson["pair"] is not None else "🔔 <b>Дополнительно</b>"
        if lesson["time"]:
            text += f"  ⏰ {lesson['time']}"
        text += f"\n📚 <b>{lesson['subject']}</b>\n"
        if lesson["teacher"]:
            text += f"👨‍🏫 {lesson['teacher']}\n"
        if lesson["room"]:
            text += f"🚪 Кабинет: {lesson['room']}\n"
        text += "\n"
    return text


def main():
    today = date.today()
    text = "🌅 <b>Доброе утро!</b>\n\n" + format_day(today, get_schedule(today))
    response = requests.post(
        f"https://api.telegram.org/bot{TOKEN}/sendMessage",
        json={"chat_id": CHAT_ID, "text": text, "parse_mode": "HTML"},
        timeout=30,
    )
    response.raise_for_status()
    print("Расписание отправлено", today)


if __name__ == "__main__":
    main()
