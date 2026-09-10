import io
import logging
import os
import re
from datetime import date, datetime, timedelta, time
from zoneinfo import ZoneInfo
from typing import Optional

import openpyxl
import requests

from fastapi import FastAPI, Request
from fastapi.responses import PlainTextResponse

from telegram import (
    Update,
    ReplyKeyboardMarkup,
)
from telegram.ext import (
    Application,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

# ============================================================
# НАСТРОЙКИ
# ============================================================

# Для компьютера можно вставить токен прямо сюда.
# Для Render токен лучше хранить в Environment Variables.
TOKEN = os.getenv("BOT_TOKEN")
AUTO_CHAT_ID = os.getenv("AUTO_CHAT_ID")

# ID твоей Google-таблицы
SPREADSHEET_ID = "1UmJKEHIy7vYc3matpjZ9ZC7NqiVY3QHETa4f4DotAlY"

# Наша группа
GROUP = "451"

# В таблице:
# K = предмет / преподаватель
# L = кабинет
SUBJECT_COL = 11
ROOM_COL = 12

# В таблице:
# A = дата
# C = номер пары
# D = время
DATE_COL = 1
PAIR_COL = 3
TIME_COL = 4

# Часовой пояс школы.Q
# Барнаул / Алтайский край = UTC+7.
TIMEZONE = "Asia/Barnaul"
LOCAL_TZ = ZoneInfo(TIMEZONE)

# Во сколько автоматически присылать расписание.
AUTO_SEND_TIME = time(hour=7, minute=0, tzinfo=LOCAL_TZ)

def today_local() -> date:
    """Текущая дата по часовому поясу школы (Барнаул)."""
    return datetime.now(LOCAL_TZ).date()


# ============================================================
# ЛОГИ
# ============================================================

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)

logger = logging.getLogger(__name__)




# ============================================================
# ОТСЛЕЖИВАНИЕ СООБЩЕНИЙ БОТА
# ============================================================

tracked_messages = {}


async def reply_tracked(update: Update, text: str, **kwargs):
    """Отправляет ответ и запоминает его, чтобы /clear мог удалить сообщения бота."""
    message = await update.message.reply_text(text, **kwargs)
    chat_id = update.effective_chat.id if update.effective_chat else None
    if chat_id is not None:
        tracked_messages.setdefault(chat_id, set()).add(message.message_id)
    return message


def track_message(chat_id: int, message_id: int):
    tracked_messages.setdefault(chat_id, set()).add(message_id)


# ============================================================
# КНОПКИ
# ============================================================

MAIN_KEYBOARD = ReplyKeyboardMarkup(
    [
        ["📅 Сегодня", "➡️ Завтра"],
        ["📆 Эта неделя", "🔄 Обновить"],
        ["⚙️ Авто-расписание"],
    ],
    resize_keyboard=True,
)


# ============================================================
# GOOGLE SHEETS
# ============================================================

def download_google_sheet():
    """
    Скачивает актуальную версию Google Sheets в формате XLSX.
    """

    url = (
        f"https://docs.google.com/spreadsheets/d/"
        f"{SPREADSHEET_ID}/export?format=xlsx"
    )

    response = requests.get(
        url,
        timeout=30,
    )

    response.raise_for_status()

    return openpyxl.load_workbook(
        io.BytesIO(response.content),
        data_only=True,
    )


# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================

def clean(value) -> Optional[str]:
    """
    Приводит значение ячейки к аккуратному тексту.
    """

    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    # Убираем лишние переводы строк
    text = re.sub(r"\s+", " ", text)

    return text


def normalize_date(value):
    """
    Преобразует содержимое ячейки в date.
    """

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    return None


def day_name(d: date):
    names = [
        "Понедельник",
        "Вторник",
        "Среда",
        "Четверг",
        "Пятница",
        "Суббота",
        "Воскресенье",
    ]

    return names[d.weekday()]


# ============================================================
# ПОЛУЧЕНИЕ РАСПИСАНИЯ
# ============================================================

def get_schedule(target_date: date):
    """
    Возвращает расписание группы 451 на указанную дату.
    """

    workbook = download_google_sheet()

    result = []

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        current_date = None

        # Ищем даты и расписание внутри листа
        for row in range(1, ws.max_row + 1):

            # ------------------------------------------------
            # Определяем текущую дату
            # ------------------------------------------------

            value = ws.cell(
                row=row,
                column=DATE_COL,
            ).value

            parsed_date = normalize_date(value)

            if parsed_date:
                current_date = parsed_date

            if current_date != target_date:
                continue

            # ------------------------------------------------
            # Читаем основные значения
            # ------------------------------------------------

            pair_value = ws.cell(
                row=row,
                column=PAIR_COL,
            ).value

            lesson_time = clean(
                ws.cell(
                    row=row,
                    column=TIME_COL,
                ).value
            )

            subject = clean(
                ws.cell(
                    row=row,
                    column=SUBJECT_COL,
                ).value
            )

            room = clean(
                ws.cell(
                    row=row,
                    column=ROOM_COL,
                ).value
            )

            # ------------------------------------------------
            # Обычная пара
            # ------------------------------------------------

            if pair_value is not None and subject:

                try:
                    pair = int(float(pair_value)) # type: ignore
                except (ValueError, TypeError):
                    pair = None

                teacher = None

                # В твоей таблице преподаватель находится
                # на следующей строке в колонке K.
                next_row = row + 1

                if next_row <= ws.max_row:

                    possible_teacher = clean(
                        ws.cell(
                            row=next_row,
                            column=SUBJECT_COL,
                        ).value
                    )

                    if possible_teacher:

                        # Если это не такой же текст,
                        # считаем это преподавателем.
                        if possible_teacher != subject:
                            teacher = possible_teacher

                result.append(
                    {
                        "pair": pair,
                        "time": lesson_time,
                        "subject": subject,
                        "teacher": teacher,
                        "room": room,
                    }
                )

            # ------------------------------------------------
            # Особые строки без номера пары
            #
            # Например:
            # "Классный час"
            # ------------------------------------------------

            elif lesson_time and subject:

                # Проверяем, не является ли значение
                # преподавателем.
                lower_subject = subject.lower()

                teacher_words = [
                    "преподаватель:",
                ]

                if not any(
                    word in lower_subject
                    for word in teacher_words
                ):

                    result.append(
                        {
                            "pair": None,
                            "time": lesson_time,
                            "subject": subject,
                            "teacher": None,
                            "room": room,
                        }
                    )

    # Убираем дубли
    unique = []

    seen = set()

    for item in result:

        key = (
            item["pair"],
            item["time"],
            item["subject"],
            item["room"],
        )

        if key not in seen:

            seen.add(key)
            unique.append(item)

    # Сортируем по номеру пары
    unique.sort(
        key=lambda x: (
            x["pair"] is None,
            x["pair"] if x["pair"] is not None else 999,
        )
    )

    return unique


# ============================================================
# ФОРМАТИРОВАНИЕ
# ============================================================

def format_day(target_date: date, lessons):
    """
    Красиво формирует сообщение Telegram.
    """

    text = (
        f"📅 <b>{day_name(target_date)}</b>\n"
        f"📆 {target_date.strftime('%d.%m.%Y')}\n"
        f"🎓 <b>Группа 451</b>\n"
        f"\n"
    )

    if not lessons:

        text += "😎 <b>Пар нет.</b> Можно отдыхать!"

        return text

    for lesson in lessons:

        if lesson["pair"] is not None:

            text += (
                f"🔔 <b>{lesson['pair']} пара</b>"
            )

        else:

            text += "🔔 <b>Дополнительно</b>"

        if lesson["time"]:
            text += f"  ⏰ {lesson['time']}"

        text += "\n"

        text += (
            f"📚 <b>{lesson['subject']}</b>\n"
        )

        if lesson["teacher"]:

            text += (
                f"👨‍🏫 {lesson['teacher']}\n"
            )

        if lesson["room"]:

            text += (
                f"🚪 Кабинет: {lesson['room']}\n"
            )

        text += "\n"

    return text


# ============================================================
# ПОЛУЧЕНИЕ СЕГОДНЯ
# ============================================================

# ============================================================
# СЕГОДНЯ
# ============================================================

async def send_today(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    try:
        target_date = today_local()
        lessons = get_schedule(target_date)
        text = format_day(target_date, lessons)

        assert update.message is not None
        await reply_tracked(update, 
            text,
            parse_mode="HTML",
            reply_markup=MAIN_KEYBOARD,
        )

    except Exception:
        logger.exception("Ошибка получения расписания на сегодня")

        assert update.message is not None
        await reply_tracked(update, 
            "❌ Не удалось получить расписание.",
            reply_markup=MAIN_KEYBOARD,
        )


async def send_tomorrow(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):

    try:

        target_date = (
            today_local()
            + timedelta(days=1)
        )

        lessons = get_schedule(target_date)

        text = format_day(
            target_date,
            lessons,
        )

        assert update.message is not None
        await reply_tracked(update, 
            text,
            parse_mode="HTML",
            reply_markup=MAIN_KEYBOARD,
        )

    except Exception:

        logger.exception(
            "Ошибка получения расписания"
        )

        assert update.message is not None
        await reply_tracked(update, 
            "❌ Не удалось получить расписание.",
            reply_markup=MAIN_KEYBOARD,
        )


# ============================================================
# НЕДЕЛЯ
# ============================================================

async def send_week(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):

    try:

        today_date = today_local()

        monday = (
            today_date
            - timedelta(
                days=today_date.weekday()
            )
        )

        full_text = (
            "🎓 <b>Расписание группы 451</b>\n"
            "📆 <b>Текущая неделя</b>\n\n"
        )

        for i in range(6):

            target_date = (
                monday
                + timedelta(days=i)
            )

            lessons = get_schedule(
                target_date
            )

            full_text += format_day(
                target_date,
                lessons,
            )

            full_text += (
                "\n"
                "━━━━━━━━━━━━━━\n\n"
            )

        # Telegram имеет ограничение на размер сообщения.
        # Отправляем большими кусками.
        for start in range(
            0,
            len(full_text),
            3900,
        ):

            assert update.message is not None
            await reply_tracked(update, 
                full_text[
                    start:start + 3900
                ],
                parse_mode="HTML",
                reply_markup=MAIN_KEYBOARD,
            )

    except Exception:

        logger.exception(
            "Ошибка получения недели"
        )

        assert update.message is not None
        await reply_tracked(update, 
            "❌ Не удалось получить расписание.",
            reply_markup=MAIN_KEYBOARD,
        )


# ============================================================
# АВТОМАТИЧЕСКАЯ ОТПРАВКА
# ============================================================

async def automatic_schedule(
    context: ContextTypes.DEFAULT_TYPE,
):

    job = context.job

    assert job is not None
    chat_id = job.chat_id

    try:

        target_date = today_local()

        lessons = get_schedule(
            target_date
        )

        text = format_day(
            target_date,
            lessons,
        )

        sent = await context.bot.send_message(
            chat_id=chat_id, # type: ignore
            text=(
                "🌅 <b>Доброе утро!</b>\n\n"
                + text
            ),
            parse_mode="HTML",
            reply_markup=MAIN_KEYBOARD,
        )
        track_message(chat_id, sent.message_id)

    except Exception:

        logger.exception(
            "Ошибка автоматической отправки"
        )


async def enable_auto(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    chat = update.effective_chat
    if chat is None:
        return

    if context.job_queue is None:
        await reply_tracked(
            update,
            "❌ Локальное авто-расписание недоступно. Проверь зависимость python-telegram-bot[job-queue].",
            reply_markup=MAIN_KEYBOARD,
        )
        return

    # Удаляем старую задачу этого чата, если она была.
    for job in context.job_queue.get_jobs_by_name(f"schedule_{chat.id}"):
        job.schedule_removal()

    context.job_queue.run_daily(
        automatic_schedule,
        time=AUTO_SEND_TIME,
        chat_id=chat.id,
        name=f"schedule_{chat.id}",
    )

    await reply_tracked(
        update,
        "✅ <b>Авто-расписание включено.</b>\n\n"
        "Каждый день в <b>07:00 по времени Барнаула</b> я буду присылать расписание в этот чат.",
        parse_mode="HTML",
        reply_markup=MAIN_KEYBOARD,
    )


async def disable_auto(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    chat = update.effective_chat
    if chat is None:
        return

    if context.job_queue is not None:
        for job in context.job_queue.get_jobs_by_name(f"schedule_{chat.id}"):
            job.schedule_removal()

    await reply_tracked(
        update,
        "🛑 <b>Авто-расписание выключено.</b>",
        parse_mode="HTML",
        reply_markup=MAIN_KEYBOARD,
    )

# ============================================================
# START
# ============================================================

async def start(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):

    assert update.message is not None
    await reply_tracked(update, 
        "👋 <b>Привет!</b>\n\n"
        "Я бот расписания группы <b>451</b>.\n\n"
        "Выбери нужный пункт ниже 👇",
        parse_mode="HTML",
        reply_markup=MAIN_KEYBOARD,
    )


# ============================================================
# ОБНОВИТЬ
# ============================================================

async def refresh(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):

    assert update.message is not None
    await reply_tracked(update, 
        "🔄 Загружаю самое свежее расписание..."
    )

    await send_today(
        update,
        context,
    )


# ============================================================
# ОБРАБОТКА КНОПОК
# ============================================================

async def button_handler(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):

    assert update.message is not None
    text = update.message.text

    if text == "📅 Сегодня":

        await send_today(
            update,
            context,
        )

    elif text == "➡️ Завтра":

        await send_tomorrow(
            update,
            context,
        )

    elif text == "📆 Эта неделя":

        await send_week(
            update,
            context,
        )

    elif text == "🔄 Обновить":

        await refresh(
            update,
            context,
        )

    elif text == "⚙️ Авто-расписание":

        await enable_auto(
            update,
            context,
        )


# ============================================================
# ПОЛУЧИТЬ CHAT ID
# ============================================================

async def myid(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    chat = update.effective_chat
    if chat is None:
        return

    await update.effective_message.reply_text(
        f"Ваш chat_id: <code>{chat.id}</code>",
        parse_mode="HTML",
    )




# ============================================================
# ПОМОЩЬ
# ============================================================

async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await reply_tracked(
        update,
        "📚 <b>Команды бота</b>\n\n"
        "/start — запустить бота\n"
        "/today — расписание на сегодня\n"
        "/tomorrow — расписание на завтра\n"
        "/week — расписание на неделю\n"
        "/auto — включить авто-расписание в этом чате\n"
        "/stop — выключить авто-расписание в этом чате\n"
        "/myid — показать chat_id\n"
        "/clear — удалить сообщения бота, которые он запомнил в этом чате\n"
        "/help — показать эту справку",
        parse_mode="HTML",
        reply_markup=MAIN_KEYBOARD,
    )


# ============================================================
# ОЧИСТКА
# ============================================================

async def clear_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    chat = update.effective_chat
    message = update.effective_message
    if chat is None or message is None:
        return

    # Удаляем саму команду пользователя. Для групп это работает, если у бота
    # есть право удалять сообщения.
    try:
        await message.delete()
    except Exception:
        pass

    message_ids = tracked_messages.pop(chat.id, set())
    for message_id in list(message_ids):
        try:
            await context.bot.delete_message(
                chat_id=chat.id,
                message_id=message_id,
            )
        except Exception:
            pass


# ============================================================
# MAIN
# ============================================================

application = Application.builder().token(TOKEN or "").build()

application.add_handler(CommandHandler("start", start))
application.add_handler(CommandHandler("myid", myid))
application.add_handler(CommandHandler("help", help_command))
application.add_handler(CommandHandler("clear", clear_command))
application.add_handler(CommandHandler("today", send_today))
application.add_handler(CommandHandler("tomorrow", send_tomorrow))
application.add_handler(CommandHandler("week", send_week))
application.add_handler(CommandHandler("auto", enable_auto))
application.add_handler(CommandHandler("stop", disable_auto))
application.add_handler(
    MessageHandler(filters.TEXT & ~filters.COMMAND, button_handler)
)

app = FastAPI()


@app.get("/", response_class=PlainTextResponse)
async def health():
    return "Schedule bot is running"


@app.api_route(
    "/health",
    methods=["GET", "HEAD"],
    response_class=PlainTextResponse,
)
async def health_check():
    """Endpoint for an external uptime monitor."""
    return "OK"

@app.post("/alice")
async def alice_webhook(request: Request):
    """Webhook для навыка Алисы."""

    try:
        data = await request.json()

        command = (
            data.get("request", {})
            .get("command", "")
            .strip()
            .lower()
        )

        # По умолчанию показываем сегодня
        target_date = today_local()

        # Если пользователь сказал "завтра"
        if "завтра" in command:
            target_date += timedelta(days=1)

        lessons = get_schedule(target_date)

        if not lessons:
            text = (
                f"На {target_date.strftime('%d.%m')} "
                "пар нет."
            )
        else:
            parts = []

            for lesson in lessons:
                pair = lesson["pair"]
                lesson_time = lesson["time"]
                subject = lesson["subject"]

                if pair is not None:
                    prefix = f"{pair} пара"
                else:
                    prefix = "Дополнительно"

                if lesson_time:
                    parts.append(
                        f"{prefix}, {lesson_time} — {subject}"
                    )
                else:
                    parts.append(
                        f"{prefix} — {subject}"
                    )

            text = (
                f"Расписание на {target_date.strftime('%d.%m')}. "
                + ". ".join(parts)
                + "."
            )

        return {
            "response": {
                "text": text,
                "end_session": True,
            },
            "version": "1.0",
        }

    except Exception:
        logger.exception("Ошибка навыка Алисы")

        return {
            "response": {
                "text": (
                    "Не удалось получить расписание. "
                    "Попробуй ещё раз."
                ),
                "end_session": True,
            },
            "version": "1.0",
        }
    


@app.post("/telegram/webhook")
async def telegram_webhook(request: Request):
    data = await request.json()
    update = Update.de_json(data, application.bot)
    await application.process_update(update)
    return {"ok": True}


@app.on_event("startup")
async def startup():
    if not TOKEN:
        raise RuntimeError("BOT_TOKEN не задан в Render Environment Variables")
    if not os.getenv("RENDER_EXTERNAL_URL"):
        raise RuntimeError("RENDER_EXTERNAL_URL не найден")

    await application.initialize()
    await application.start()

    webhook_url = os.getenv("RENDER_EXTERNAL_URL").rstrip("/") + "/telegram/webhook"
    await application.bot.set_webhook(url=webhook_url)
    logger.info("Webhook установлен: %s", webhook_url)


@app.on_event("shutdown")
async def shutdown():
    await application.stop()
    await application.shutdown()


# ============================================================
